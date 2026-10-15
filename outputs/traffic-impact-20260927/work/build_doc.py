from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path(__file__).resolve().parent.parent
d=Document()
sec=d.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
for style in d.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(21)
d.styles['Title'].paragraph_format.space_after=Pt(6)
d.styles['Heading 1'].font.size=Pt(12)
d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
d.styles['Heading 1'].paragraph_format.space_after=Pt(6)
d.core_properties.title='Assignment 3 Traffic Impact Analysis'
d.core_properties.author='Leo Leung'
def p(t,bold=False):
 after_table=len(d.element.body)>1 and d.element.body[-2].tag==qn('w:tbl')
 x=d.add_paragraph();r=x.add_run(t);r.bold=bold
 if after_table:x.paragraph_format.space_before=Pt(6)
 return x
def h(t):d.add_heading(t,level=1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths): t.columns[i].width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 for vals in rows:
  for cell,v in zip(t.add_row().cells,vals):cell.text=str(v)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  no=OxmlElement('w:cantSplit');trpr.append(no)
  if ri==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EBF0' if ri==0 else ('F6F8FA' if ri%2 else 'FFFFFF'));pr.append(sh)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.0
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if (ri==0 or ci<2 or headers[ci].endswith('LOS')) else WD_ALIGN_PARAGRAPH.RIGHT
    for r in para.runs:r.font.size=Pt(9);r.bold=(ri==0)
 return t

d.add_paragraph('Assignment 3 Traffic Impact Analysis','Title')
p('Leo Leung | CPLN 550 Introduction to Transportation Planning | Fall 2026')
p('The proposed 800,000-square-foot office development increases peak-direction demand on all three highways. The Southwest corridor has the largest speed losses, while NE in the AM and NW in the PM move from LOS D to E.')
h('1 Current conditions')
p('I match both the observed flow per lane and speed in Appendix B to Appendix A. Speed identifies the correct demand when the same observed flow occurs under different congestion conditions. All demand and flow figures below are vehicles per hour across all lanes.')
table(['Highway','Period','Lanes','Demand','Observed\nflow','Speed\n(mph)','LOS'],[
 ['NE','AM',6,'10,800','10,800',52,'D'],['SW','AM',4,'10,000','6,000',35,'F'],['NW','AM',5,'8,000','8,000',54,'D'],
 ['NE','PM',6,'10,200','10,200',53,'D'],['SW','PM',4,'9,600','6,400',39,'F'],['NW','PM',5,'9,000','9,000',52,'D']], [.75,.65,.6,1.4,1.4,1.2,1.1])
p('For SW in the AM, 6,000 / 4 = 1,500 vehicles per lane per hour at 35 mph corresponds to demand of 2,500 per lane, or 10,000 total. In the PM, 6,400 / 4 = 1,600 at 39 mph corresponds to demand of 2,400 per lane, or 9,600 total. Demand exceeds observed flow by 4,000 vehicles/hour in the AM and 3,200 in the PM. On NE and NW, existing demand equals observed flow.')
h('2 Office trip generation')
p('The office contains 800,000 / 1,000 = 800 trip-generation units. AM trips in both directions are 800 × 1.45 = 1,160. Multiplying by 87% gives 1,009.20 inbound trips in the AM peak hour. PM trips in both directions are 800 × 1.50 = 1,200. Multiplying by 85% gives 1,020.00 outbound trips in the PM peak hour.')
h('3 New trips by corridor')
p('Each corridor receives its Appendix C share of peak-direction trips. For example, SW receives 1,009.20 × 40% = 403.68 AM trips. I retain fractional trip estimates until the required per-lane demand lookup.')
table(['Highway','AM share','AM inbound trips','PM share','PM outbound trips'],[
 ['NE','35%','353.22','30%','306.00'],['SW','40%','403.68','45%','459.00'],['NW','25%','252.30','25%','255.00'],['Total','100%','1,009.20','100%','1,020.00']], [.85,1.05,2.0,1.05,2.15])

d.add_page_break()
h('4 Conditions after development')
p('For each corridor and period, new demand equals existing demand plus assigned office trips. I divide by the lane count, round to the nearest 100 vehicles per lane per hour, and use that row of Appendix A to obtain flow, speed, and LOS. Total flow equals the table flow per lane multiplied by the lane count.')
table(['Highway','Period','Lanes','Total\ndemand','Demand\nper lane','Rounded\nper lane'],[
 ['NE','AM',6,'11,153.22','1,858.87','1,900'],['SW','AM',4,'10,403.68','2,600.92','2,600'],['NW','AM',5,'8,252.30','1,650.46','1,700'],
 ['NE','PM',6,'10,506.00','1,751.00','1,800'],['SW','PM',4,'10,059.00','2,514.75','2,500'],['NW','PM',5,'9,255.00','1,851.00','1,900']], [.85,.7,.6,1.7,1.6,1.65])
p('Updated Appendix B. Demand and flow are total vehicles/hour; speeds are mph.')
table(['Highway','Lanes','AM\ndemand','AM\nflow','AM\nspeed','AM\nLOS','PM\ndemand','PM\nflow','PM\nspeed','PM\nLOS'],[
 ['NE',6,'11,153.22','11,400',51,'E','10,506.00','10,800',52,'D'],['SW',4,'10,403.68','5,600',32,'F','10,059.00','6,000',35,'F'],['NW',5,'8,252.30','8,500',53,'D','9,255.00','9,500',51,'E']], [.72,.6,1.1,.72,.6,.48,1.1,.73,.56,.49])
p('For example, SW AM demand is 10,000 + 403.68 = 10,403.68. Dividing by four gives 2,600.92, which rounds to 2,600. Appendix A gives flow of 1,400 per lane, so total flow is 5,600 at 32 mph and LOS F. Because the assignment uses a rounded lookup, estimated flow on NE and NW can slightly exceed unrounded demand. I retain the actual demand for subsequent calculations.')
h('5 Changes in LOS and other conditions')
p('NE in the AM changes from LOS D to E, and NW in the PM changes from LOS D to E. Conditions also worsen in all four cases where the LOS letter stays unchanged:')
table(['Highway and period','LOS','Speed before and after','Flow before and after'],[
 ['SW AM','F to F','35 to 32 mph','6,000 to 5,600'],['SW PM','F to F','39 to 35 mph','6,400 to 6,000'],['NE PM','D to D','53 to 52 mph','10,200 to 10,800'],['NW AM','D to D','54 to 53 mph','8,000 to 8,500']], [1.5,.8,2.3,2.5])
h('6 Summary of development impacts')
p('The development adds 1,009.20 AM inbound trips and 1,020.00 PM outbound trips. SW receives the most new trips in each peak and is most affected: AM demand rises from 10,000 to 10,403.68 vehicles/hour while flow falls from 6,000 to 5,600 and speed falls from 35 to 32 mph. PM demand rises from 9,600 to 10,059 while flow falls from 6,400 to 6,000 and speed falls from 39 to 35 mph. SW remains LOS F in both peaks, so its unchanged grade conceals worsening congestion. Its PM speed loss of 4 mph is the largest. NE AM and NW PM worsen from D to E. NE PM and NW AM remain D but each slows by 1 mph. Flow increases on NE and NW as demand rises, while declining flow on SW reflects more severe congestion rather than fewer trips being demanded.')

d.add_page_break()
h('7 Adding one lane to the Southwest highway')
p('The added lane increases SW from four to five peak-direction lanes. Total post-development demand stays at 10,403.68 vehicles/hour in the AM and 10,059 in the PM because the scenario assumes no induced travel.')
table(['Period','Demand\nper lane','Rounded\nper lane','Flow\nper lane','Total\nflow','Speed\n(mph)','LOS'],[
 ['AM','2,080.736','2,100','1,900','9,500','48','F'],['PM','2,011.800','2,000','2,000','10,000','50','E']], [.65,1.4,1.15,1.0,1.2,1.05,.65])
p('AM: 10,403.68 / 5 = 2,080.736, rounded to 2,100. Total flow becomes 1,900 × 5 = 9,500 vehicles/hour at 48 mph. PM: 10,059 / 5 = 2,011.8, rounded to 2,000. Total flow becomes 2,000 × 5 = 10,000 vehicles/hour at 50 mph. Relative to the development without expansion, AM flow rises by 3,900 and speed by 16 mph; PM flow rises by 4,000 and speed by 15 mph. AM remains F but improves substantially, and PM improves from F to E.')
h('8 Daily vehicle hours saved')
p('Travel time per trip is distance divided by speed. I use post-development demand, including trips that cannot pass through the corridor within the observed peak hour, because all trips eventually reach their destinations. For each period, daily vehicle-hours saved equal hourly demand × 2 peak hours × (8 / old speed − 8 / new speed).')
table(['Period','Before\n(min/trip)','After\n(min/trip)','Saved\n(min/trip)','Daily trips\nfrom demand','Vehicle hours\nsaved per day'],[
 ['AM','15.0000','10.0000','5.0000','20,807.36','1,733.9467'],['PM','13.7143','9.6000','4.1143','20,118.00','1,379.5200']], [.65,1.1,1.1,1.1,1.55,1.6])
p('AM savings = 10,403.68 × 2 × (8 / 32 − 8 / 48) = 1,733.9467 vehicle-hours/day.\nPM savings = 10,059 × 2 × (8 / 35 − 8 / 50) = 1,379.5200 vehicle-hours/day.')
p('Total daily savings = 3,113.4667 vehicle-hours, or approximately 3,113.47 vehicle-hours/day.',True)
h('9 Annual monetary value')
p('Annual value = daily vehicle-hours saved × $30/vehicle-hour × 350 days/year. Using the unrounded daily savings gives $32,691,400 per year, or approximately $32.69 million.')
h('10 Break even construction cost per lane mile')
p('Question 10 refers to its own answer. I interpret this as referring to the annual savings calculated in Question 9. Ten years of savings equal $32,691,400 × 10 = $326,914,000. Adding one lane over 8 miles creates 8 lane-miles.')
p('Break-even construction cost = $326,914,000 / 8 = $40,864,250 per lane-mile, or approximately $40.86 million per lane-mile.',True)
p('This is the assignment’s undiscounted ten-year comparison. It does not subtract maintenance costs or apply a discount rate.')
p('Source: Assignment 3 Project evaluation part I, CPLN 550, Fall 2026, scenario, Questions 1–10, and Appendices A–C. Calculations retain full precision except for the required Appendix A lookup rounding.')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
d.save(out/'Traffic Impact Written Answers.docx')
print(out/'Traffic Impact Written Answers.docx')
